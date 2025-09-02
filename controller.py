# controller.py - Main orchestrator and public API for report generation
import pandas as pd
from datetime import datetime
import os
from openpyxl.styles import Alignment
import tempfile

# Import the new modular components - FIXED PATHS
from data_processors import read_rid_file_from_stream, read_metrics_file_from_stream, apply_pid_observation_logic
from excel_generators import (
    add_pivot_and_format, add_check_results_pivot, 
    create_denylist_draft_sheet, apply_denylist_conditional_formatting,
    apply_combined_data_formatting
)

def generate_survey_report(
    rid_file_stream, metrics_file_stream, survey_loi_mapping, output_dir,
    conversion_rate_threshold=10,
    security_terms_threshold=30,
    speeder_multiplier=3,
    high_loi_multiplier=3,
    negative_recs_rate_threshold=15,
    use_datetime_for_newuser=True  # Kept for compatibility but ignored
):
    """
    Processes the survey files and generates an Excel report with observations and a pivot table.
    Note: use_datetime_for_newuser is ignored - always uses date-only comparison now.
    survey_loi_mapping: Dictionary mapping survey IDs to their respective LOI values
    """
    try:
        # Validate survey_loi_mapping
        if not survey_loi_mapping or not isinstance(survey_loi_mapping, dict):
            raise ValueError("Survey LOI mapping is required and must be a dictionary of survey_id: loi_value pairs.")
        
        # Validate LOI values
        for survey_id, loi_value in survey_loi_mapping.items():
            if not (3 <= loi_value <= 100):
                raise ValueError(f"Survey LOI for {survey_id} must be between 3 and 100, got {loi_value}.")

        # Read files with enhanced error handling
        try:
            rid_df = read_rid_file_from_stream(rid_file_stream)
        except Exception as e:
            raise ValueError(f"RID file error: {str(e)}")
            
        try:
            metrics_df = read_metrics_file_from_stream(metrics_file_stream)
        except Exception as e:
            raise ValueError(f"PID Metrics file error: {str(e)}")

        # Debug: Print shapes and columns of input DataFrames
        print("DEBUG: rid_df shape:", rid_df.shape)
        print("DEBUG: rid_df columns:", rid_df.columns.tolist())
        print("DEBUG: metrics_df shape:", metrics_df.shape)
        print("DEBUG: metrics_df columns:", metrics_df.columns.tolist())

        if 'pid' not in rid_df.columns or 'pid' not in metrics_df.columns:
            raise ValueError("Critical Error: 'pid' column not found in one or both input files. Please check that files contain 'pid' column.")
        
        # Enhanced merge validation - PIDs are already strings from file readers
        rid_pids = rid_df['pid']
        metrics_pids = metrics_df['pid']
        pid_intersection = set(rid_pids) & set(metrics_pids)
        print("DEBUG: Number of PIDs in RID file:", len(rid_pids))
        print("DEBUG: Number of PIDs in Metrics file:", len(metrics_pids))
        print("DEBUG: Number of matching PIDs:", len(pid_intersection))

        if len(pid_intersection) == 0:
            raise ValueError("No matching PIDs found between RID and Metrics files. Please ensure both files contain the same PIDs.")

        # PIDs are already converted to strings in the file readers
        merged_df = pd.merge(
            rid_df,
            metrics_df,
            on='pid',
            how='left',
            indicator=False
        )

        # Debug: Print merged DataFrame shape and columns after merge
        print("DEBUG: merged_df shape:", merged_df.shape)
        print("DEBUG: merged_df columns:", merged_df.columns.tolist())

        # Add this check immediately after merge
        if merged_df.empty:
            raise ValueError(
                "Merge resulted in an empty DataFrame. "
                "No matching rows found between RID and PID files after merge. "
                "Check that your files have overlapping PIDs and correct formats."
            )

        # Scale system_conversion_rate and net_recs_rate to 0-100 immediately after merging
        if "system_conversion_rate" in merged_df.columns:
            merged_df["system_conversion_rate"] = merged_df["system_conversion_rate"] * 100
        
        # Scale NET RECS RATE using exact column name
        if 'net recs rate' in merged_df.columns:
            merged_df['net recs rate'] = merged_df['net recs rate'] * 100

        # Apply observation logic and add check columns
        merged_df = apply_pid_observation_logic(
            merged_df,
            survey_loi_mapping=survey_loi_mapping,
            conversion_rate_threshold=conversion_rate_threshold,
            security_terms_threshold=security_terms_threshold,
            speeder_multiplier=speeder_multiplier,
            high_loi_multiplier=high_loi_multiplier,
            negative_recs_rate_threshold=negative_recs_rate_threshold,
            session_loi_checks=True,
            use_datetime_for_newuser=False  # Always date-only now
        )

        # --- Insert Security Terms Rate column using exact column name ---
        if 'security terms rate' in merged_df.columns:
            # Rename to standardized name
            merged_df = merged_df.rename(columns={'security terms rate': 'Security_Terms_Rate'})
            
            # Move Security_Terms_Rate to after system_conversion_rate
            cols = list(merged_df.columns)
            if "system_conversion_rate" in cols and "Security_Terms_Rate" in cols:
                security_rate_values = merged_df["Security_Terms_Rate"]
                merged_df.drop("Security_Terms_Rate", axis=1, inplace=True)
                cols = list(merged_df.columns)
                idx = cols.index("system_conversion_rate") + 1
                merged_df.insert(idx, "Security_Terms_Rate", security_rate_values)

        # --- Insert first/last entry date match column using exact column names ---
        if 'first_entry_date' in merged_df.columns and 'last_entry_date' in merged_df.columns:
            match_col = (merged_df['first_entry_date'] == merged_df['last_entry_date'])
            cols = list(merged_df.columns)
            if 'last_entry_date' in cols:
                idx = cols.index('last_entry_date') + 1
                merged_df.insert(idx, "FirstLastDateMatch", match_col)
            else:
                merged_df["FirstLastDateMatch"] = match_col

        # --- Insert Diff Days column after last_entry_date using exact column names ---
        if 'first_entry_date' in merged_df.columns and 'last_entry_date' in merged_df.columns:
            # Convert to datetime if not already
            merged_df["first_entry_date_dt"] = pd.to_datetime(merged_df['first_entry_date'], errors='coerce')
            merged_df["last_entry_date_dt"] = pd.to_datetime(merged_df['last_entry_date'], errors='coerce')
            diff_days = (merged_df["last_entry_date_dt"] - merged_df["first_entry_date_dt"]).dt.days
            # Insert after last_entry_date
            cols = list(merged_df.columns)
            if 'last_entry_date' in cols:
                idx = cols.index('last_entry_date') + 1
                merged_df.insert(idx, "Diff Days", diff_days)
            else:
                merged_df["Diff Days"] = diff_days
            # Remove temp columns
            merged_df.drop(["first_entry_date_dt", "last_entry_date_dt"], axis=1, inplace=True)
        else:
            merged_df["Diff Days"] = None

        # --- Round columns to 2 decimal places using exact column names ---
        if "system_conversion_rate" in merged_df.columns:
            merged_df["system_conversion_rate"] = merged_df["system_conversion_rate"].round(2)
        if 'net recs rate' in merged_df.columns:
            merged_df['net recs rate'] = merged_df['net recs rate'].round(2)
        if "Security_Terms_Rate" in merged_df.columns:
            merged_df["Security_Terms_Rate"] = merged_df["Security_Terms_Rate"].round(2)

        # --- Add Flag_Count column at the end ---
        flag_columns = [
            "Poor_Conv_Rate", "New_User_Bot", "High_Security",
            "Speeder", "High_LOI", "High_RR"
        ]
        merged_df["Flag_Count"] = merged_df[flag_columns].sum(axis=1)

        # Remove blank columns (all values are NaN or empty) before writing to Excel
        merged_df = merged_df.dropna(axis=1, how='all')

        # --- Generate Excel File ---
        os.makedirs(output_dir, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_filename = f"RID-PID_Report_{timestamp}.xlsx"
        output_path = os.path.join(output_dir, output_filename)

        try:
            print("DEBUG: Starting Excel file creation...")
            with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
                # Convert data types before writing to ensure Excel compatibility
                merged_df_copy = merged_df.copy()
                
                # Ensure numeric columns are properly typed
                numeric_cols = ['system_conversion_rate', 'Security_Terms_Rate', 'Flag_Count', 'Diff Days']
                if 'net recs rate' in merged_df_copy.columns:
                    numeric_cols.append('net recs rate')
                
                for col in numeric_cols:
                    if col in merged_df_copy.columns:
                        merged_df_copy[col] = pd.to_numeric(merged_df_copy[col], errors='coerce')
                
                # Ensure text columns are strings
                text_cols = ['pid', 'supplier_bu', 'Observation']
                for col in text_cols:
                    if col in merged_df_copy.columns:
                        merged_df_copy[col] = merged_df_copy[col].astype(str).replace('nan', '')
                
                # Ensure boolean columns are properly formatted
                bool_cols = ['Poor_Conv_Rate', 'New_User_Bot', 'High_Security', 'Speeder', 'High_LOI', 'High_RR', 'FirstLastDateMatch']
                for col in bool_cols:
                    if col in merged_df_copy.columns:
                        merged_df_copy[col] = merged_df_copy[col].astype(bool)
                
                # Convert all remaining object columns to strings to avoid XML issues
                for col in merged_df_copy.columns:
                    if merged_df_copy[col].dtype == 'object' and col not in bool_cols:
                        merged_df_copy[col] = merged_df_copy[col].astype(str).replace('nan', '')
                
                print("DEBUG: Writing Combined Data sheet...")
                merged_df_copy.to_excel(writer, sheet_name='Combined Data', index=False)
                print("DEBUG: Combined Data sheet written successfully")
                
                print("DEBUG: Calling add_pivot_and_format...")
                add_pivot_and_format(writer, merged_df)
                print("DEBUG: add_pivot_and_format completed successfully")
                
                print("DEBUG: Calling add_check_results_pivot...")
                add_check_results_pivot(writer, merged_df)
                print("DEBUG: add_check_results_pivot completed successfully")

                print("DEBUG: Reordering sheets...")
                # Reorder sheets to match desired sequence
                wb = writer.book
                desired_order = [
                    "Combined Data",
                    "Flags Pivot (Priority)", 
                    "Flags Pivot (Multi)",
                    "Pivot EntryDate x Supplier",
                    "Pivot EntryDate x Flags"
                ]
                
                # Reorder existing sheets
                existing_sheets = []
                for sheet_name in desired_order:
                    if sheet_name in wb.sheetnames:
                        existing_sheets.append(wb[sheet_name])
                
                # Remove all sheets from workbook
                wb._sheets.clear()
                
                # Add sheets back in desired order
                for sheet in existing_sheets:
                    wb._sheets.append(sheet)
                print("DEBUG: Sheet reordering completed")

                print("DEBUG: Creating DenyList_Draft sheet...")
                # --- Create DenyList_Draft sheet ---
                deny_sheet, deny_df = create_denylist_draft_sheet(wb, merged_df)
                print("DEBUG: DenyList_Draft sheet created")

                # Apply conditional formatting for DenyList_Draft
                apply_denylist_conditional_formatting(deny_sheet, deny_df, merged_df)
                print("DEBUG: DenyList_Draft conditional formatting completed")
                
                print("DEBUG: Starting header alignment and data type setting...")
                # Set header alignment to left for Combined Data and DenyList_Draft
                left_align = Alignment(horizontal='left')
                # Combined Data
                combined_sheet = wb["Combined Data"]
                
                for cell in combined_sheet[1]:
                    cell.alignment = left_align
                    cell.data_type = 's'  # Headers should be text
                    
                # DenyList_Draft
                for cell in deny_sheet[1]:
                    cell.alignment = left_align
                print("DEBUG: Header alignment and data types completed")

                print("DEBUG: Starting conditional formatting for Combined Data...")
                # --- Enhanced Conditional Formatting for Combined Data ---
                apply_combined_data_formatting(combined_sheet, merged_df)

                print("DEBUG: Excel file creation completed successfully")

        except PermissionError:
            raise ValueError(f"Cannot write to output file. Please ensure the file is not open in Excel: {output_filename}")
        except Exception as e:
            print(f"DEBUG: Error in Excel file creation: {e}")
            import traceback
            traceback.print_exc()
            if "openpyxl" in str(e):
                raise ValueError(f"Excel generation error: {str(e)}. The data was processed but some visual formatting may be missing.")
            raise ValueError(f"Error creating Excel report: {str(e)}")

        return str(output_path)
        
    except ValueError:
        print("DEBUG: Caught ValueError, re-raising...")
        raise  # Re-raise ValueError as-is
    except Exception as e:
        print(f"DEBUG: Caught unexpected exception: {e}")
        import traceback
        traceback.print_exc()
        raise ValueError(f"Unexpected error during report generation: {str(e)}")

def generate_pid_only_report(
    metrics_file_stream,
    output_dir,
    conversion_rate_threshold=10,
    security_terms_threshold=30,
    negative_recs_rate_threshold=15,
    use_datetime_for_newuser=True
):
    """Processes only the PID Metrics file and generates an Excel report with observations (PID-only mode)."""
    try:
        # Enhanced file reading with error handling
        try:
            metrics_df = read_metrics_file_from_stream(metrics_file_stream)
        except Exception as e:
            raise ValueError(f"PID Metrics file error: {str(e)}")
            
        if metrics_df.empty:
            raise ValueError("PID Metrics file contains no data rows. Please check your Excel file.")
        
        # Pass None for speeder_multiplier and high_loi_multiplier, and session_loi_checks=False
        metrics_df = apply_pid_observation_logic(
            metrics_df,
            survey_loi_mapping=None,
            conversion_rate_threshold=conversion_rate_threshold,
            security_terms_threshold=security_terms_threshold,
            speeder_multiplier=None,
            high_loi_multiplier=None,
            negative_recs_rate_threshold=negative_recs_rate_threshold,
            session_loi_checks=False,
            use_datetime_for_newuser=use_datetime_for_newuser
        )
        
        # --- Generate Excel File ---
        os.makedirs(output_dir, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_filename = f"pid_metrics_report_{timestamp}.xlsx"
        output_path = os.path.join(output_dir, output_filename)
        
        print("DEBUG: PID-only output_path:", repr(output_path))
        
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            metrics_df.to_excel(writer, sheet_name='PID Metrics Data', index=False)
            workbook = writer.book
            ws_data = writer.sheets['PID Metrics Data']
            ws_data.freeze_panes = ws_data['B2']
            ws_data.auto_filter.ref = ws_data.dimensions
            
            # Add Observation Pivot with formatting and conditional totals
            if 'Observation' in metrics_df.columns:
                obs_counts = metrics_df['Observation'].value_counts().reset_index()
                obs_counts.columns = ['Observation', 'Count']
                na_row = obs_counts[obs_counts['Observation'] == '-n/a-']
                other_rows = obs_counts[obs_counts['Observation'] != '-n/a-'].sort_values('Count', ascending=False)
                obs_counts_sorted = pd.concat([na_row, other_rows], ignore_index=True)
                
                obs_counts_sorted.to_excel(writer, sheet_name='Observation Pivot', index=False)
                ws_pivot = writer.sheets['Observation Pivot']
                ws_pivot.auto_filter.ref = ws_pivot.dimensions
                ws_pivot.freeze_panes = ws_pivot['B2']
                
                from openpyxl.styles import Font
                dark_green_font = Font(color="006400")
                for row in ws_pivot.iter_rows(min_row=2, max_row=2, min_col=1, max_col=2):
                    for cell in row:
                        if cell.value == '-n/a-' or (cell.row == 2 and ws_pivot['A2'].value == '-n/a-'):
                            cell.font = dark_green_font
                
                # Add column total
                ws_pivot.cell(row=ws_pivot.max_row+1, column=1, value='Column Total')
                ws_pivot.cell(row=ws_pivot.max_row, column=2, value=obs_counts_sorted['Count'].sum())
                ws_pivot.cell(row=ws_pivot.max_row, column=2).font = Font(bold=True)
                    
        return str(output_path)
    except ValueError:
        raise  # Re-raise ValueError as-is
    except Exception as e:
        print(f"DEBUG: PID-only unexpected error: {e}")
        raise ValueError(f"Unexpected error in PID-only processing: {str(e)}")